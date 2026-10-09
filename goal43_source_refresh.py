"""Adapt newly converted Office files through the existing metro_data readers."""

import argparse
import calendar
from datetime import date, datetime, time
from decimal import Decimal
from html import escape
import hashlib
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl import load_workbook

from metro_data.pipeline import (
    METRICS, aggregate_stations, load_map, parse_quarter_sheet, read_hourly,
    read_transport, require,
)


class QuarterSheet:
    def __init__(self, worksheet):
        self.worksheet, self.title = worksheet, worksheet.title

    def iter_rows(self, values_only=True):
        for index, original in enumerate(self.worksheet.iter_rows(values_only=True), 1):
            row = list(original)
            if index == 2:
                for column in range(2, min(98, len(row))):
                    if isinstance(row[column], str):
                        row[column] = time.fromisoformat(row[column])
            yield row


def read_converted_quarter(path, mapping, source_id):
    book = load_workbook(path, read_only=True, data_only=True)
    tables, checks, empty = [], [], []
    try:
        for worksheet in book:
            if worksheet.title == "Лист1":
                require(not any(value is not None for row in worksheet.iter_rows(values_only=True) for value in row),
                        "Лист1 must remain empty")
                empty.append(worksheet.title)
                continue
            table, audit = parse_quarter_sheet(QuarterSheet(worksheet), mapping, source_id)
            tables.append(table)
            checks.extend(audit)
    finally:
        book.close()
    require(tables, "no service day sheets")
    result = pd.concat(tables, ignore_index=True).sort_values(
        ["service_date", "vestibule_id", "interval_start"]).reset_index(drop=True)
    days = sorted(result.service_date.unique())
    require(len(days) == len(tables), "duplicate service dates")
    require(len({(day.year, day.month) for day in days}) == 1, "multiple months")
    expected = [date(days[0].year, days[0].month, day)
                for day in range(1, calendar.monthrange(days[0].year, days[0].month)[1] + 1)]
    return result, checks, {"days": len(days), "rows": len(result), "empty_sheets": empty,
                           "missing_days": [str(day) for day in expected if day not in days],
                           "adapter": "only ISO time strings in C2:CT2 converted to datetime.time"}


def read_converted_transport(path, source_id, url):
    """Reuse HTML transport validation after retaining native XLSX cell locations."""
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        require(len(book.sheetnames) == 1, "transport workbook must have one sheet")
        worksheet = book.active
        physical = list(worksheet.iter_rows(values_only=True))
        require([str(int(v)) for v in physical[5][1:7]] == ["1", "2", "3", "4", "5", "6"],
                "transport line columns changed")
        html_rows, locations = [], {}
        for index, row in enumerate(physical, 1):
            if not row or row[0] is None:
                continue
            if row[0] in METRICS:
                require(len(row) >= 8 and all(value is not None for value in row[:8]), "missing transport metric cell")
                cells = row[:8]
                require(row[0] not in locations, "duplicate transport metric")
                locations[row[0]] = f"{worksheet.title}!B{index}"
            else:
                cells = [row[0]]
            html_rows.append("<tr>" + "".join("<td>" + escape(str(value)) + "</td>" for value in cells) + "</tr>")
        html_rows.insert(0, "<tr>" + "".join(f"<td>{line}</td>" for line in range(1, 7)) + "</tr>")
        with TemporaryDirectory(prefix="goal43-transport-") as folder:
            adapted = Path(folder) / "transport.html"
            adapted.write_text('<html><meta charset="utf-8"><table>' + "".join(html_rows) + "</table></html>", encoding="utf-8")
            result, checks = read_transport(adapted, source_id, url)
        result["source_path"] = path.as_posix()
        result["source_locator"] = result.original_label.map(locations)
        return result, checks
    finally:
        book.close()


def semantic_comparison(old, new):
    keys = ["vestibule_id", "interval_start"]
    if any(table.duplicated(keys).any() for table in (old, new)):
        raise ValueError("duplicate entry identity")
    joined = old[keys + ["entry_count"]].merge(new[keys + ["entry_count"]], on=keys, how="outer",
                                             suffixes=("_old", "_new"), indicator=True, validate="one_to_one")
    both = joined._merge.eq("both")
    equal = (joined.entry_count_old.eq(joined.entry_count_new) |
             (joined.entry_count_old.isna() & joined.entry_count_new.isna())).fillna(False)
    changed = int((both & ~equal).sum())
    added, removed = int(joined._merge.eq("right_only").sum()), int(joined._merge.eq("left_only").sum())
    return {"legacy_rows": len(old), "new_rows": len(new), "added_keys": added, "removed_keys": removed,
            "value_rows_differing": changed, "semantic_equal": not (added or removed or changed),
            "legacy_total_entries": int(old.entry_count.sum()), "new_total_entries": int(new.entry_count.sum()),
            "legacy_missing_values": int(old.entry_count.isna().sum()), "new_missing_values": int(new.entry_count.isna().sum())}


def refresh(root, raw, output):
    root = Path(root).resolve()
    raw, output = Path(raw), Path(output)
    raw = (raw if raw.is_absolute() else root / raw).resolve()
    output = (output if output.is_absolute() else root / output).resolve()
    protected = [root / path for path in ("data/raw", "docs/reference", "data/processed/metro", "data/processed/metro_repeat")]
    if output == root or any(output == path or path in output.parents for path in protected):
        raise ValueError("source refresh must use a separate output directory")
    output.mkdir(parents=True, exist_ok=True)
    mapping = load_map(root / "docs/reference/vestibule_map.csv")
    aliases = pd.read_csv(root / "docs/reference/hourly_vestibule_aliases.csv")
    new_quarter, quarter_checks, quarter_info = read_converted_quarter(
        raw / "1zNYr89aNuY60LEIwVRXYB6UXn1UA0lSD.xlsx", mapping, "D02_goal43")
    new_hourly, hourly_info = read_hourly(raw / "10cr6-19pPlJRIL4S4R3VOININJwBt-8l.xlsx", mapping, aliases, "D05_goal43")
    old_quarter = pd.read_parquet(root / "data/processed/metro/entries_vestibule_15m.parquet")
    old_quarter = old_quarter.loc[old_quarter.source_id.eq("D02")].copy()
    old_hourly = pd.read_parquet(root / "data/processed/metro/entries_vestibule_1h.parquet")
    station = aggregate_stations(new_quarter, mapping)
    old_stations = pd.read_parquet(root / "data/processed/metro/entries_station_15m.parquet")
    entries_all = pd.concat([old_stations.loc[old_stations.source_id.ne("D02")], station], ignore_index=True)
    entries_all = entries_all.sort_values(["service_date", "station_id", "interval_start"]).reset_index(drop=True)
    for name, table in (("entries_vestibule_september_15m", new_quarter), ("entries_station_september_15m", station),
                        ("entries_station_15m", entries_all), ("entries_vestibule_1h", new_hourly)):
        table.to_parquet(output / f"{name}.parquet", index=False)
    transport, transport_checks = read_converted_transport(
        raw / "1RH2kw2e2NLvxa6byMo82F4-h1eD1UiXa.xlsx", "D08_goal43",
        "https://drive.google.com/file/d/1RH2kw2e2NLvxa6byMo82F4-h1eD1UiXa/view")
    transport.to_csv(output / "transport_work_q2.csv", index=False, lineterminator="\n")
    old_transport = pd.read_csv(root / "data/processed/metro/transport_work.csv", dtype={"value": str})
    old_transport = old_transport.loc[old_transport.source_id.eq("D08")]
    joined = old_transport[["metric", "value"]].merge(transport[["metric", "value"]], on="metric", suffixes=("_old", "_new"), validate="one_to_one")
    transport_equal = all(Decimal(a) == Decimal(b)
                          for a, b in zip(joined.value_old, joined.value_new)) and len(joined) == len(old_transport) == len(transport)
    report = {"retrieved_at": datetime.now(ZoneInfo("Europe/Moscow")).isoformat(),
              "D02": {**semantic_comparison(old_quarter, new_quarter), "extraction": quarter_info},
              "D05": {**semantic_comparison(old_hourly, new_hourly), "extraction": hourly_info},
              "D08": {"legacy_format": "HTML stored as .xls", "new_format": "xlsx",
                      "period": [str(transport.period_start.iloc[0]), str(transport.period_end.iloc[0])],
                      "line_id": 1, "metric_rows": len(transport), "semantic_equal": transport_equal,
                      "metrics": transport.to_dict(orient="records")},
              "checks": {"quarter": quarter_checks, "transport": transport_checks},
              "inputs": {path.as_posix(): hashlib.sha256(path.read_bytes()).hexdigest() for path in sorted(raw.glob("*.xlsx"))},
              "legacy_outputs_overwritten": False}
    (output / "semantic_comparison.json").write_text(json.dumps(report, ensure_ascii=False, default=str, indent=2) + "\n", encoding="utf-8")
    return {name: {key: value for key, value in report[name].items() if key not in {"extraction", "metrics"}} for name in ("D02", "D05", "D08")}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", default=Path(__file__).resolve().parent, type=Path)
    parser.add_argument("--raw", default="data/raw/metro/goal4_3_source_versions/20261008", type=Path)
    parser.add_argument("--output", default="data/processed/goal4_3_source_refresh", type=Path)
    arguments = parser.parse_args()
    print(json.dumps(refresh(arguments.root, arguments.raw, arguments.output), indent=2))
