from datetime import date, time, timedelta
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from metro_data.pipeline import (
    SourceError, aggregate_stations, interval_start, load_map,
    parse_quarter_sheet, read_hourly, read_transport, sniff_format,
)
from metro_data.build import check, validate_tables
from metro_data.pipeline import read_quarter

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def mapping():
    return load_map(ROOT / 'docs/reference/vestibule_map.csv')


@pytest.fixture
def sheet(mapping):
    ws = Workbook().active
    ws.title = 'ТВхП 28022026 С 15-мин.'
    ws.cell(1, 1, 'Таблица входных потоков. 28022026 Суббота')
    for i in range(96):
        t = interval_start(date(2026, 2, 28), i)
        ws.cell(2, i + 3, t.time())
        ws.cell(3, i + 3, 24)
    for i, label in enumerate(mapping['quarter_hour_label'], 4):
        ws.cell(i, 1, label)
        ws.cell(i, 2, 96)
        for j in range(3, 99):
            ws.cell(i, j, 1)
    ws.cell(28, 2, 24 * 96)
    return ws


def test_midnight_and_month_boundary():
    assert interval_start(date(2026, 2, 28), 84).isoformat() == '2026-03-01T00:00:00+03:00'
    last = interval_start(date(2026, 2, 28), 95)
    assert (last + timedelta(minutes=15)).isoformat() == '2026-03-01T03:00:00+03:00'
    assert str(last.tzinfo) == 'Europe/Moscow'


def test_totals_are_checks_not_observations(sheet, mapping):
    data, checks = parse_quarter_sheet(sheet, mapping, 'fixture')
    assert len(data) == 24 * 96
    assert int(data.entry_count.sum()) == 24 * 96
    assert len(checks) == 24 + 96 + 2
    assert all(c['status'] == 'pass' for c in checks)


def test_extra_line_daily_total_in_cu3_is_checked_and_excluded(sheet, mapping):
    sheet.cell(3, 99, 2305)
    data, checks = parse_quarter_sheet(sheet, mapping, 'fixture')
    assert len(data) == 2304 and data.entry_count.sum() == 2304
    c = next(c for c in checks if c['locator'] == 'CU3')
    assert c['status'] == 'mismatch' and c['delta'] == -1


def test_empty_sheet_is_skipped_but_new_content_fails(tmp_path, sheet, mapping):
    empty = sheet.parent.create_sheet('Лист1')
    path = tmp_path / 'quarter.xlsx'
    sheet.parent.save(path)
    data, checks, info = read_quarter(path, mapping, 'fixture')
    assert len(data) == 2304 and info['empty_sheets'] == ['Лист1']
    empty.cell(1, 1, 'Новые данные')
    sheet.parent.save(path)
    with pytest.raises(SourceError, match='no longer empty'):
        read_quarter(path, mapping, 'fixture')


def test_station_provenance_cannot_silently_mix_sources(sheet, mapping):
    data, _ = parse_quarter_sheet(sheet, mapping, 'fixture')
    data.loc[0, 'source_id'] = 'another_file'
    with pytest.raises(SourceError, match='provenance'):
        aggregate_stations(data, mapping)


def test_check_notices_new_source_files(tmp_path):
    raw = tmp_path / 'data/raw/metro'
    raw.mkdir(parents=True)
    (raw / 'new.txt').write_text('new source')
    output = tmp_path / 'out'
    output.mkdir()
    (output / 'build_manifest.json').write_text('{"inputs": {}, "outputs": {}}')
    with pytest.raises(SourceError, match='Source file set changed'):
        check(tmp_path, output)


def test_missing_is_not_zero_and_station_is_incomplete(sheet, mapping):
    sheet.cell(4, 3).value = None
    sheet.cell(5, 3).value = 0
    data, checks = parse_quarter_sheet(sheet, mapping, 'fixture')
    assert pd.isna(data.loc[data.vestibule_id == 'V01', 'entry_count'].iloc[0])
    assert data.loc[data.vestibule_id == 'V02', 'entry_count'].iloc[0] == 0
    station = aggregate_stations(data, mapping)
    row = station[station.station_id == 'L1_01'].iloc[0]
    assert row.vestibules_present == 1 and row.vestibules_expected == 2
    assert row.entry_count == 0 and not row.is_complete
    assert any(c['status'] == 'incomplete' for c in checks)
    # An entirely absent station stays null when the expected grid is restored.
    data = data[data.station_id != 'L1_03']
    absent = aggregate_stations(data, mapping).query("station_id == 'L1_03'")
    assert absent.entry_count.isna().all() and not absent.is_complete.any()


@pytest.mark.parametrize('change', ['label', 'time', 'extra', 'date', 'negative', 'fraction'])
def test_changed_structure_or_invalid_count_fails(sheet, mapping, change):
    if change == 'label': sheet.cell(4, 1, 'Неизвестный вход')
    if change == 'time': sheet.cell(2, 3, time(3, 15))
    if change == 'extra': sheet.cell(29, 3, 1)
    if change == 'date': sheet.cell(1, 1, '29022026')
    if change == 'negative': sheet.cell(4, 3, -1)
    if change == 'fraction': sheet.cell(4, 3, 0.5)
    with pytest.raises(SourceError):
        parse_quarter_sheet(sheet, mapping, 'fixture')


def test_daily_discrepancy_is_retained(sheet, mapping):
    sheet.cell(4, 2, 97)
    data, checks = parse_quarter_sheet(sheet, mapping, 'fixture')
    assert data.entry_count.sum() == 2304
    c = next(c for c in checks if c['check'] == 'vestibule_daily' and c['locator'] == 'B4')
    assert c['status'] == 'mismatch' and c['delta'] == -1


@pytest.mark.parametrize('damage', ['duplicate', 'time', 'station_sum'])
def test_saved_data_validation_detects_corruption(sheet, mapping, damage):
    v, _ = parse_quarter_sheet(sheet, mapping, 'fixture')
    s = aggregate_stations(v, mapping)
    h = v.iloc[:1].copy()
    h['interval_end'] = h.interval_start + pd.Timedelta(hours=1)
    if damage == 'duplicate': v = pd.concat([v, v.iloc[:1]])
    if damage == 'time': v.loc[0, 'interval_start'] += pd.Timedelta(minutes=15)
    if damage == 'station_sum': s.loc[0, 'entry_count'] += 1
    with pytest.raises(SourceError):
        validate_tables(v, s, h, mapping)


def test_hourly_has_no_padding_or_division(tmp_path, mapping):
    w = Workbook()
    s = w.active
    for j, label in enumerate(['Дата', 'Станция', 'Количество пассажиров'], 1):
        s.cell(7, j, label)
    for r, count in [(8, 8), (9, None), (10, 0)]:
        s.cell(r, 1, f'30.09.2026 {r - 8:02}')
        s.cell(r, 2, 'пр. Ветеранов-1')
        s.cell(r, 3).value = count
    s.cell(11, 1, 'Итого: 8')
    p = tmp_path / 'hour.xlsx'
    w.save(p)
    aliases = pd.read_csv(ROOT / 'docs/reference/hourly_vestibule_aliases.csv')
    data, info = read_hourly(p, mapping, aliases, 'fixture')
    assert len(data) == 3 and list(data.entry_count.dropna()) == [8, 0]
    assert data.entry_count.isna().sum() == 1
    assert data.interval_start.iloc[0].isoformat() == '2026-09-30T00:00:00+03:00'
    assert data.service_date.iloc[0] == date(2026, 9, 29)
    assert info['footer_reconciliation']['status'] == 'incomplete'


def test_html_xls_uses_line_one_and_preserves_discrepancy(tmp_path):
    p = tmp_path / 'MO1.xls'
    labels = [
        ('Максимальная парность (пар)', '32'), ('Поездо/км (п/км)', '100'),
        ('Общие ваг/км (в/км)', '850'), ('с пассажирами (в/км)', '800'),
        ('нулевой пробег(в/км) , в том числе', '50'),
        ('нулевой пробег депо (в/км)', '10'), ('нулевой пробег обкатки (в/км)', '15'),
        ('Поездо/час (п/час)', '10'), ('Поездо/час в движении (п/час)', '8'),
        ('Простой (п/час)', '2,03'),
    ]
    html = '<html><meta charset="KOI8-R"><table><tr><td>за II квартал 2026 года</td></tr>'
    html += '<tr>' + ''.join(f'<td>{n}</td>' for n in range(1, 7)) + '</tr>'
    for label, value in labels:
        html += f'<tr><td>{label}</td><td>{value}</td>' + '<td>999</td>' * 6 + '</tr>'
    p.write_bytes((html + '</table></html>').encode('koi8-r'))
    assert sniff_format(p.read_bytes()) == 'html'
    data, checks = read_transport(p, 'fixture', 'fixture-url')
    assert data.loc[data.metric == 'car_km_total', 'value'].iloc[0] == '850'
    assert data.loc[data.metric == 'car_km_zero', 'raw_value'].iloc[0] == '50'
    assert data.line_id.eq(1).all()
    assert next(c for c in checks if c['check'] == 'car_km_sum')['status'] == 'pass'
    c = next(c for c in checks if c['check'] == 'train_hours_sum')
    assert c['delta'] == '0.03' and c['status'] == 'mismatch'
