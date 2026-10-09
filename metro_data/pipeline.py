from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from pathlib import Path
import calendar
import math
import re
from zoneinfo import ZoneInfo
from zipfile import is_zipfile, ZipFile
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

TZ = ZoneInfo('Europe/Moscow')
TIME_KEYS = ['service_date', 'interval_start', 'interval_end']
ENTRY_COLUMNS = ['service_date', 'vestibule_id', 'station_id', 'interval_start',
                 'interval_end', 'entry_count', 'source_id', 'source_sheet', 'quality_flag']


class SourceError(ValueError):
    pass


def require(condition, message):
    if not condition:
        raise SourceError(message)


def interval_start(day, index):
    return datetime.combine(day, time(3), TZ) + timedelta(minutes=15 * index)


def load_map(path):
    mapping = pd.read_csv(path, dtype=str)
    require(len(mapping) == 24 and mapping.local_station_key.nunique() == 19,
            'vestibule_map must contain 24 vestibules and 19 stations')
    require(mapping.quarter_hour_label.is_unique and not mapping.isna().any().any(),
            'Duplicate or missing map labels')
    require(mapping.groupby('local_station_key').station_name.nunique().eq(1).all(),
            'Conflicting station names')
    mapping['station_id'] = mapping.local_station_key
    mapping['vestibule_id'] = [f'V{i:02}' for i in range(1, 25)]
    return mapping


def count(value, locator):
    if value is None:
        return None
    require(isinstance(value, (int, float)) and not isinstance(value, bool)
            and math.isfinite(value) and value >= 0 and value == int(value), f'Invalid entry count at {locator}: {value!r}')
    return int(value)


def frame(records):
    data = pd.DataFrame(records, columns=ENTRY_COLUMNS)
    data['entry_count'] = pd.array(data.entry_count, dtype='Int64')
    return data.sort_values(['service_date', 'vestibule_id', 'interval_start']).reset_index(drop=True)


def comparison(source, sheet, name, locator, expected, actual):
    delta = None if actual is None or expected is None else actual - expected
    return dict(source_id=source, source_sheet=sheet, check=name, locator=locator,
                expected=expected, actual=actual, delta=delta,
                status='incomplete' if delta is None else 'pass' if delta == 0 else 'mismatch')


def parse_quarter_sheet(ws, mapping, source_id):
    # Read the physical extent too: populated extra cells signal a changed schema.
    rows = list(ws.iter_rows(values_only=True))
    require(len(rows) >= 28, f'{ws.title}: fewer than 28 rows')
    right_total = rows[2][98] if len(rows[2]) > 98 else None
    for r, row in enumerate(rows, 1):
        outside = row[99:] if r == 3 else row[98:] if r <= 28 else row
        require(not any(v is not None for v in outside),
                f'{ws.title}: populated cells outside A1:CT28 and the optional total CU3')
    rows = [list(row[:98]) + [None] * max(0, 98 - len(row)) for row in rows[:28]]
    header = re.search(r'\b(\d{8})\b', str(rows[0][0]))
    sheet_date = re.search(r'\b(\d{8})\b', ws.title)
    require(header and sheet_date and header[1] == sheet_date[1], f'{ws.title}: date headers differ')
    try:
        day = datetime.strptime(header[1], '%d%m%Y').date()
    except ValueError as e:
        raise SourceError(f'{ws.title}: invalid date') from e
    starts = [interval_start(day, i) for i in range(96)]
    require(rows[1][2:98] == [t.time() for t in starts], f'{ws.title}: expected 96 intervals from 03:00')
    require(rows[2][0] is None and rows[27][0] is None and
            all(v is None for v in rows[27][2:]), f'{ws.title}: total rows changed')
    labels = [row[0] for row in rows[3:27]]
    require(len(set(labels)) == 24 and set(labels) == set(mapping.quarter_hour_label),
            f'{ws.title}: vestibule labels differ from explicit map')
    lookup = mapping.set_index('quarter_hour_label')
    records, checks, values, daily = [], [], [], []
    for r, row in enumerate(rows[3:27], 4):
        m = lookup.loc[row[0]]
        entries = [count(v, f'{ws.title}:R{r}C{i+3}') for i, v in enumerate(row[2:98])]
        total = count(row[1], f'{ws.title}:B{r}')
        daily.append(total)
        values.append(entries)
        checks.append(comparison(source_id, ws.title, 'vestibule_daily', f'B{r}', total,
                                 sum(entries) if None not in entries else None))
        for start, n in zip(starts, entries):
            records.append((day, m.vestibule_id, m.station_id, start,
                            start + timedelta(minutes=15), n, source_id, ws.title,
                            'missing' if n is None else 'ok'))
    for i in range(96):
        col = [row[i] for row in values]
        total = count(rows[2][i+2], f'{ws.title}:R3C{i+3}')
        checks.append(comparison(source_id, ws.title, 'line_interval', f'R3C{i+3}', total,
                                 sum(col) if None not in col else None))
    grand = count(rows[27][1], f'{ws.title}:B28')
    flat = [n for row in values for n in row]
    checks.append(comparison(source_id, ws.title, 'line_daily', 'B28', grand,
                             sum(flat) if None not in flat else None))
    checks.append(comparison(source_id, ws.title, 'daily_totals_sum', 'B28', grand,
                             sum(daily) if None not in daily else None))
    if right_total is not None:
        checks.append(comparison(source_id, ws.title, 'line_daily_right', 'CU3',
                                 count(right_total, f'{ws.title}:CU3'), sum(flat) if None not in flat else None))
    return frame(records), checks


def read_quarter(path, mapping, source_id):
    book = load_workbook(path, read_only=True, data_only=True)
    tables, checks, empty = [], [], []
    try:
        for ws in book:
            if ws.title == 'Лист1':
                require(not any(v is not None for row in ws.iter_rows(values_only=True) for v in row),
                        f'{path.name}: Лист1 is no longer empty')
                empty.append(ws.title)
                continue
            table, audit = parse_quarter_sheet(ws, mapping, source_id)
            tables.append(table)
            checks.extend(audit)
    finally:
        book.close()
    require(tables, f'{path.name}: no day sheets')
    data = pd.concat(tables, ignore_index=True).sort_values(
        ['service_date', 'vestibule_id', 'interval_start']).reset_index(drop=True)
    dates = sorted(data.service_date.unique())
    require(len(dates) == len(tables), f'{path.name}: duplicate service dates')
    months = {(d.year, d.month) for d in dates}
    require(len(months) == 1, f'{path.name}: multiple months')
    y, m = next(iter(months))
    expected = [date(y, m, d) for d in range(1, calendar.monthrange(y, m)[1]+1)]
    return data, checks, dict(source_id=source_id, month=f'{y}-{m:02}', rows=len(data),
                             days=len(dates), missing_days=[str(d) for d in expected if d not in dates],
                             empty_sheets=empty)


def aggregate_stations(data, mapping):
    keys = TIME_KEYS + ['station_id']
    group = data.groupby(keys, observed=True)
    result = group.entry_count.sum(min_count=1).to_frame()
    result['vestibules_present'] = group.entry_count.count()
    result['vestibules_observed'] = group.vestibule_id.nunique()
    for col in ['source_id', 'source_sheet']:
        require(group[col].nunique().le(1).all(), f'Conflicting station provenance: {col}')
        result[col] = group[col].first()
    grid = data[TIME_KEYS].drop_duplicates().merge(mapping[['station_id']].drop_duplicates(), how='cross')
    result = result.reindex(pd.MultiIndex.from_frame(grid[keys])).reset_index()
    expected = mapping.groupby('station_id').size()
    result['vestibules_expected'] = result.station_id.map(expected)
    for col in ['vestibules_present', 'vestibules_observed']:
        result[col] = result[col].fillna(0).astype('int64')
    result['entry_count'] = result.entry_count.astype('Int64')
    result['is_complete'] = result.vestibules_present.eq(result.vestibules_expected)
    return result.sort_values(['service_date', 'station_id', 'interval_start']).reset_index(drop=True)


def read_hourly(path, mapping, aliases, source_id):
    require(aliases.hourly_label.is_unique and aliases.quarter_hour_label.isin(mapping.quarter_hour_label).all(),
            'Invalid explicit hourly aliases')
    lookup = aliases.merge(mapping, on='quarter_hour_label', validate='many_to_one').set_index('hourly_label')
    book = load_workbook(path, read_only=True, data_only=True)
    records, excluded = [], []
    try:
        require(len(book.sheetnames) == 1, f'{path.name}: expected one hourly sheet')
        ws = book.active
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i < 7:
                continue
            if i == 7:
                require(list(row[:3]) == ['Дата', 'Станция', 'Количество пассажиров'], 'Hourly header changed')
                continue
            if all(v is None for v in row):
                continue
            if str(row[0]).startswith(('Общее количество пассажиров за выбранный период:', 'Итого:')):
                require(all(v is None for v in row[1:]), f'{path.name}: hourly footer changed')
                excluded.append(dict(row=i, text=row[0]))
                continue
            require(len(row) >= 3 and not any(v is not None for v in row[3:]), f'Hourly row {i} changed')
            try:
                start = datetime.strptime(str(row[0]), '%d.%m.%Y %H').replace(tzinfo=TZ)
            except ValueError as e:
                raise SourceError(f'Hourly date changed at row {i}: {row[0]!r}') from e
            require(row[1] in lookup.index, f'Unknown hourly vestibule: {row[1]!r}')
            m, n = lookup.loc[row[1]], count(row[2], f'hourly:R{i}C3')
            records.append(((start - timedelta(hours=3)).date(), m.vestibule_id, m.station_id,
                            start, start + timedelta(hours=1), n, source_id, ws.title,
                            'missing' if n is None else 'ok'))
    finally:
        book.close()
    require(records, 'No hourly records')
    data = frame(records)
    require(len(excluded) == 1, 'Hourly total footer changed or missing')
    total = re.search(r'(\d+)\s*$', excluded[0]['text'])
    require(total, 'Hourly total footer value changed')
    footer = comparison(source_id, ws.title, 'hourly_total', f'A{excluded[0]["row"]}', int(total[1]),
                        int(data.entry_count.sum()) if data.entry_count.notna().all() else None)
    return data, dict(excluded_rows=excluded, source_id=source_id, footer_reconciliation=footer)


def sniff_format(raw):
    if raw.startswith(b'%PDF-'): return 'pdf'
    if raw.startswith(b'\x89PNG\r\n\x1a\n'): return 'png'
    if raw.startswith(b'\xff\xd8\xff'): return 'jpeg'
    if raw.startswith(b'\xd0\xcf\x11\xe0'): return 'ole_excel'
    if is_zipfile(BytesIO(raw)):
        with ZipFile(BytesIO(raw)) as z:
            names = set(z.namelist())
        if 'xl/workbook.xml' in names: return 'xlsx'
        if 'word/document.xml' in names: return 'docx'
        if 'ppt/presentation.xml' in names: return 'pptx'
        return 'zip'
    if re.search(br'<(?:html|table)\b', raw[:4000], re.I): return 'html'
    if b'\x00' not in raw: return 'text'
    return 'unknown'


class HTMLRows(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows, self.row, self.cell = [], [], None

    def handle_starttag(self, tag, attrs):
        if tag == 'tr': self.row = []
        if tag in ('td', 'th'): self.cell = []

    def handle_data(self, data):
        if self.cell is not None: self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in ('td', 'th') and self.cell is not None:
            self.row.append(' '.join(''.join(self.cell).split()))
            self.cell = None
        if tag == 'tr': self.rows.append(self.row)


METRICS = {
    'Максимальная парность (пар)': ('max_pairs', 'pairs/h'),
    'Поездо/км (п/км)': ('train_km', 'train_km'),
    'Общие ваг/км (в/км)': ('car_km_total', 'car_km'),
    'с пассажирами (в/км)': ('car_km_passenger', 'car_km'),
    'нулевой пробег(в/км) , в том числе': ('car_km_zero', 'car_km'),
    'нулевой пробег депо (в/км)': ('car_km_depot', 'car_km'),
    'нулевой пробег обкатки (в/км)': ('car_km_test', 'car_km'),
    'Поездо/час (п/час)': ('train_hours', 'train_h'),
    'Поездо/час в движении (п/час)': ('moving_hours', 'train_h'),
    'Простой (п/час)': ('idle_hours', 'train_h'),
}


def read_transport(path, source_id, url):
    raw = path.read_bytes()
    require(sniff_format(raw) == 'html', f'{path.name}: expected HTML, detected {sniff_format(raw)}')
    encoding = re.search(br'charset\s*=\s*["\']?([\w-]+)', raw, re.I)
    require(encoding, f'{path.name}: missing declared encoding')
    try:
        text = raw.decode(encoding[1].decode('ascii'))
    except (LookupError, UnicodeError) as e:
        raise SourceError(f'{path.name}: invalid encoding') from e
    parser = HTMLRows()
    parser.feed(text)
    rows = parser.rows
    require(['1', '2', '3', '4', '5', '6'] in rows, f'{path.name}: line columns changed')
    period = next((r[0] for r in rows if len(r) == 1 and re.search(r'за .*20\d\d', r[0])), '')
    match = re.search(r'за (?:(I{1,3}|IV) квартал )?(20\d\d)', period)
    require(match, f'{path.name}: unsupported period')
    year, quarter = int(match[2]), match[1]
    first = 1 if not quarter else 1 + 3 * (['I', 'II', 'III', 'IV'].index(quarter))
    last = first + 2 if quarter else 12
    start, end = date(year, first, 1), date(year, last, calendar.monthrange(year, last)[1])
    values, records = {}, []
    for i, row in enumerate(rows, 1):
        if not row or row[0] not in METRICS: continue
        require(len(row) == 8, f'{path.name}: metric columns changed at row {i}')
        metric, unit = METRICS[row[0]]
        require(metric not in values, f'{path.name}: duplicate metric {metric}')
        try:
            n = Decimal(row[1].replace(' ', '').replace('\xa0', '').replace(',', '.'))
        except InvalidOperation as e:
            raise SourceError(f'{path.name}: invalid metric value') from e
        require(n.is_finite() and n >= 0, f'{path.name}: invalid metric value')
        values[metric] = n
        records.append(dict(period_start=start, period_end=end, period_label=period, line_id=1,
                            metric=metric, original_label=row[0], raw_value=row[1], value=str(n),
                            unit=unit, source_id=source_id, source_path=path.as_posix(), source_url=url,
                            source_locator=f'HTML tr[{i}]/td[2]', quality_flag='source_observed'))
    require(len(values) == len(METRICS), f'{path.name}: required metric missing or renamed')
    checks = []
    for name, actual, expected in [
        ('car_km_sum', values['car_km_passenger'] + values['car_km_zero'], values['car_km_total']),
        ('train_hours_sum', values['moving_hours'] + values['idle_hours'], values['train_hours']),
        ('eight_cars_vs_passenger_km', values['train_km'] * 8, values['car_km_passenger']),
    ]:
        c = comparison(source_id, '', name, period, expected, actual)
        checks.append({k: str(v) if isinstance(v, Decimal) else v for k, v in c.items()})
    require(values['car_km_depot'] + values['car_km_test'] <= values['car_km_zero'],
            f'{path.name}: zero mileage subsets exceed total')
    return pd.DataFrame(records), checks
